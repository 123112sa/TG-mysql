from __future__ import annotations

import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from webapp.modules.schedule.repository import get_schedule_plan


REPO_ROOT = Path(__file__).resolve().parents[4]
REPORT_DIR = Path(os.getenv("REPORT_DIR", str(REPO_ROOT / "reports")))
EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REPORT_EXTENSION = ".xlsx"

SCHEDULE_REPORT_FIELDS = [
    {"key": "product", "label": "硬质糖果产量", "unit": "t"},
    {"key": "jb", "label": "JB投入量", "unit": "L/min"},
    {"key": "mxx", "label": "MXX投入量", "unit": "L/min"},
    {"key": "ldx", "label": "L/DX投入量", "unit": "L/min"},
    {"key": "dnx", "label": "DNX投入量", "unit": "L/min"},
    {"key": "sa", "label": "SA投入量", "unit": "L/min"},
    {"key": "fsa", "label": "FSA投入量", "unit": "L/min"},
    {"key": "ylsn", "label": "YLSN投入量", "unit": "L/min"},
    {"key": "nx", "label": "NX投入量", "unit": "L/min"},
    {"key": "mxxflc", "label": "MXXFLC产量", "unit": "L/min"},
    {"key": "x2_2d", "label": "X2_2D产量", "unit": "L/min"},
    {"key": "cztg", "label": "CZTG产量", "unit": "L/min"},
    {"key": "txxhcc", "label": "TXXHCC产量", "unit": "L/min"},
    {"key": "yttg", "label": "YTTG产量", "unit": "L/min"},
    {"key": "tg", "label": "TG产量", "unit": "t"},
]

INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

HEADER_FILL = PatternFill("solid", fgColor="9BB0D5")
TITLE_FILL = PatternFill("solid", fgColor="2F5FAF")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
WHITE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(color="1F3557", bold=True)
NORMAL_FONT = Font(color="1F3557")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)


def _format_header_label(header: str) -> str:
    return re.sub(r"\s*\((L/min|t)\)$", r"\n(\1)", str(header))


def _preferred_column_width(header: str) -> float:
    header_text = str(header)
    if "L/min" in header_text:
        return 18.0
    if "(t)" in header_text or "产量" in header_text:
        return 16.0
    if "日期" in header_text:
        return 14.0
    if "星期" in header_text:
        return 10.0
    if "累计" in header_text or "进度" in header_text:
        return 13.0
    return 12.0


def _safe_filename(plan_name: str) -> str:
    cleaned = INVALID_FILENAME_CHARS.sub("_", str(plan_name or "").strip())
    cleaned = cleaned.strip(" .") or "schedule_report"
    return f"{cleaned}.xlsx"


def _iso_from_timestamp(timestamp: float) -> str:
    return datetime.fromtimestamp(timestamp, tz=timezone.utc).isoformat()


def _report_path(file_name: str) -> Path:
    normalized_name = str(file_name or "").strip()
    if (
        not normalized_name
        or Path(normalized_name).name != normalized_name
        or not normalized_name.lower().endswith(REPORT_EXTENSION)
        or INVALID_FILENAME_CHARS.search(normalized_name)
    ):
        raise ValueError("报表文件名非法")

    report_dir = REPORT_DIR.resolve()
    path = (report_dir / normalized_name).resolve()
    if report_dir not in path.parents:
        raise ValueError("报表文件路径非法")
    return path


def list_report_files() -> list[dict]:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    reports = []

    for path in sorted(
        REPORT_DIR.glob(f"*{REPORT_EXTENSION}"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    ):
        if not path.is_file():
            continue
        stat = path.stat()
        reports.append(
            {
                "reportName": path.stem,
                "fileName": path.name,
                "updatedAt": _iso_from_timestamp(stat.st_mtime),
                "size": stat.st_size,
            }
        )

    return reports


def get_existing_report_file(file_name: str) -> tuple[Path, str]:
    path = _report_path(file_name)
    if not path.is_file():
        raise FileNotFoundError("报表文件不存在")
    return path, path.name


def _number(value) -> float:
    try:
        return round(float(value or 0), 3)
    except (TypeError, ValueError):
        return 0.0


def _append_title(ws, title: str, column_count: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    cell = ws.cell(row=1, column=1, value=title)
    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def _append_metadata(ws, metadata: Iterable[tuple[str, object]], start_row: int = 3) -> int:
    row_index = start_row
    for label, value in metadata:
        ws.cell(row=row_index, column=1, value=label)
        ws.cell(row=row_index, column=2, value=value)
        ws.cell(row=row_index, column=1).fill = SECTION_FILL
        for col in range(1, 3):
            cell = ws.cell(row=row_index, column=col)
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
        row_index += 1
    return row_index


def _append_table(ws, headers: list[str], rows: list[list[object]], start_row: int) -> None:
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=column_index, value=_format_header_label(header))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        column_letter = get_column_letter(column_index)
        ws.column_dimensions[column_letter].width = _preferred_column_width(header)

    ws.row_dimensions[start_row].height = 34

    for row_offset, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(row=start_row + row_offset, column=column_index, value=value)
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(value, (int, float)):
                cell.number_format = "0.000"

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _auto_fit_columns(ws) -> None:
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            parts = value.splitlines() or [value]
            max_length = max(max_length, *(len(part) for part in parts))
        current_width = ws.column_dimensions[column_letter].width or 0
        ws.column_dimensions[column_letter].width = min(max(current_width, max_length + 4, 12), 24)


def _style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal=cell.alignment.horizontal or "center", vertical="center")


def _daily_summary_rows(daily_results: list[dict], target_amount: float) -> list[list[object]]:
    rows = []
    cumulative = 0.0
    for item in daily_results:
        product = _number(item.get("product"))
        cumulative += product
        progress = min(cumulative / target_amount, 1.0) if target_amount > 0 else 0.0
        rows.append(
            [
                item.get("date", ""),
                item.get("dayName", ""),
                _number(item.get("jb")),
                _number(item.get("mxx")),
                _number(item.get("ldx")),
                _number(item.get("dnx")),
                _number(item.get("sa")),
                _number(item.get("fsa")),
                _number(item.get("ylsn")),
                _number(item.get("nx")),
                product,
                round(progress, 4),
            ]
        )
    return rows


def _build_process_rows(daily_results: list[dict], fields: list[str]) -> list[list[object]]:
    rows = []
    for item in daily_results:
        rows.append([item.get("date", ""), *[_number(item.get(field)) for field in fields]])
    return rows


def _set_progress_format(ws, header_row: int, column_index: int) -> None:
    for row_index in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=row_index, column=column_index).number_format = "0.0%"


def generate_schedule_report(schedule: dict) -> dict:
    plan = schedule.get("plan") or {}
    daily_results = schedule.get("dailyResults") or []
    if not daily_results:
        raise ValueError("没有可生成报表的每日优化结果")

    plan_name = str(plan.get("planName") or f"{plan.get('startDate', '')}_{plan.get('planningDays', '')}d").strip()
    planning_days = int(plan.get("planningDays") or len(daily_results))
    target_amount = _number(plan.get("targetAmount"))

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    report_path = REPORT_DIR / _safe_filename(plan_name)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = f"{planning_days}日计划汇总"

    summary_headers = [
        "日期",
        "星期",
        "JB投入量 (L/min)",
        "MXX投入量 (L/min)",
        "L/DX投入量 (L/min)",
        "DNX投入量 (L/min)",
        "SA投入量 (L/min)",
        "FSA投入量 (L/min)",
        "YLSN投入量 (L/min)",
        "NX投入量 (L/min)",
        "硬质糖果产量 (t)",
        "累计进度",
    ]
    _append_title(summary_ws, f"{planning_days}日计划汇总", len(summary_headers))
    next_row = _append_metadata(
        summary_ws,
        [
            ("方案名称", plan_name),
            ("计划生产总量 (t)", target_amount),
            ("生产起止日期", f"{plan.get('startDate', '')} 至 {plan.get('endDate', '')}"),
            ("生产周期 (天)", planning_days),
            ("路线", plan.get("route", "")),
        ],
    )
    summary_header_row = next_row + 1
    _append_table(summary_ws, summary_headers, _daily_summary_rows(daily_results, target_amount), summary_header_row)
    _set_progress_format(summary_ws, summary_header_row, len(summary_headers))

    process_specs = [
        ("XH1D", ["日期", "JB投入量 (L/min)", "MXXFLC产量 (L/min)"], ["jb", "mxxflc"]),
        ("XH23D", ["日期", "MXXFLC投入量 (L/min)", "MXX投入量 (L/min)", "L/DX投入量 (L/min)", "X2_2D产量 (L/min)", "CZTG产量 (L/min)"], ["mxxflc", "mxx", "ldx", "x2_2d", "cztg"]),
        ("YX", ["日期", "CZTG投入量 (L/min)", "TXXHCC产量 (L/min)"], ["cztg", "txxhcc"]),
        ("JZ", ["日期", "TXXHCC投入量 (L/min)", "YLSN投入量 (L/min)", "YTTG产量 (L/min)"], ["txxhcc", "ylsn", "yttg"]),
        ("ZP", ["日期", "YTTG投入量 (L/min)", "TG产量 (t)"], ["yttg", "tg"]),
    ]

    for sheet_name, headers, fields in process_specs:
        ws = wb.create_sheet(sheet_name)
        _append_title(ws, sheet_name, len(headers))
        _append_table(ws, headers, _build_process_rows(daily_results, fields), 3)

    for ws in wb.worksheets:
        _style_sheet(ws)
        _auto_fit_columns(ws)

    wb.save(report_path)
    return {
        "fileName": report_path.name,
        "downloadUrl": f"/api/reports/schedule/{plan.get('id')}",
    }


def get_schedule_report_file(plan_id: int) -> tuple[Path, str]:
    schedule = get_schedule_plan(plan_id)
    if not schedule:
        raise ValueError("未找到对应的历史方案")

    report_info = generate_schedule_report(schedule)
    report_path = REPORT_DIR / report_info["fileName"]
    return report_path, report_info["fileName"]


def get_schedule_report_fields() -> list[dict]:
    return [dict(item) for item in SCHEDULE_REPORT_FIELDS]


def _schedule_report_field_map() -> dict[str, dict]:
    return {str(item["key"]): item for item in SCHEDULE_REPORT_FIELDS}


def _normalize_selected_fields(selected_fields: Iterable[str]) -> list[dict]:
    field_map = _schedule_report_field_map()
    normalized = []
    seen = set()

    for field in selected_fields or []:
        key = str(field or "").strip()
        if not key or key in seen:
            continue
        if key not in field_map:
            raise ValueError(f"未知的报表变量：{key}")
        normalized.append(field_map[key])
        seen.add(key)

    if not normalized:
        raise ValueError("请至少选择一个报表变量")
    return normalized


def _field_header(field: dict) -> str:
    unit = str(field.get("unit") or "").strip()
    label = str(field.get("label") or field.get("key") or "")
    return f"{label} ({unit})" if unit else label


def _custom_report_filename(plan_name: str) -> str:
    base_name = _safe_filename(plan_name)[:-len(REPORT_EXTENSION)]
    return f"{base_name}_自定义报表{REPORT_EXTENSION}"


def generate_custom_schedule_report(schedule: dict, selected_fields: Iterable[str]) -> dict:
    plan = schedule.get("plan") or {}
    daily_results = schedule.get("dailyResults") or []
    if not daily_results:
        raise ValueError("没有可生成报表的每日优化结果")

    fields = _normalize_selected_fields(selected_fields)
    plan_name = str(plan.get("planName") or f"{plan.get('startDate', '')}_{plan.get('planningDays', '')}d").strip()
    planning_days = int(plan.get("planningDays") or len(daily_results))
    target_amount = _number(plan.get("targetAmount"))

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    report_path = REPORT_DIR / _custom_report_filename(plan_name)

    headers = ["日期", "星期", *[_field_header(field) for field in fields]]
    rows = [
        [
            item.get("date", ""),
            item.get("dayName", ""),
            *[_number(item.get(field["key"])) for field in fields],
        ]
        for item in daily_results
    ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "自定义报表"
    _append_title(worksheet, f"{planning_days}日自定义报表", len(headers))
    next_row = _append_metadata(
        worksheet,
        [
            ("方案名称", plan_name),
            ("计划生产总量 (t)", target_amount),
            ("生产起止日期", f"{plan.get('startDate', '')} 至 {plan.get('endDate', '')}"),
            ("生产周期 (天)", planning_days),
            ("路线", plan.get("route", "")),
            ("选择变量", "、".join(str(field.get("label")) for field in fields)),
        ],
    )
    _append_table(worksheet, headers, rows, next_row + 1)
    _style_sheet(worksheet)
    _auto_fit_columns(worksheet)

    workbook.save(report_path)
    return {
        "fileName": report_path.name,
        "downloadUrl": f"/api/reports/schedule/{plan.get('id')}/custom",
    }


def get_custom_schedule_report_file(plan_id: int, selected_fields: Iterable[str]) -> tuple[Path, str]:
    schedule = get_schedule_plan(plan_id)
    if not schedule:
        raise ValueError("未找到对应的历史方案")

    report_info = generate_custom_schedule_report(schedule, selected_fields)
    report_path = REPORT_DIR / report_info["fileName"]
    return report_path, report_info["fileName"]
